from __future__ import annotations

import json
from pathlib import Path

from ..safety import resolve_workspace_path
from .registry import json_result, registry


def _extract_docx_text(path: Path) -> str:
    try:
        from docx import Document

        doc = Document(str(path))
        paragraphs = [p.text for p in doc.paragraphs]
        table_texts = []
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells]
                table_texts.append(" | ".join(row_texts))
        parts = []
        if paragraphs:
            parts.append("\n".join(paragraphs))
        if table_texts:
            parts.append("\n--- TABLES ---\n" + "\n".join(table_texts))
        return "\n".join(parts)
    except ImportError:
        return "[ERROR: python-docx not installed. Install with: pip install python-docx]"
    except Exception as exc:
        return f"[ERROR extracting .docx content: {exc}]"


def _read_file(args: dict, runtime: dict) -> str:
    raw_path = args.get("path")
    max_chars = int(args.get("max_chars") or 20000)
    offset = max(0, int(args.get("offset") or 0))
    try:
        path = resolve_workspace_path(raw_path)
    except ValueError:
        path = Path(raw_path).expanduser().resolve()
    if not path.is_file():
        return json_result(success=False, error=f"File not found: {path}")
    if path.suffix.lower() == ".docx":
        text = _extract_docx_text(path)
    else:
        text = path.read_text(encoding="utf-8", errors="replace")
    total_chars = len(text)
    window = text[offset:]
    content = window[:max_chars]
    has_more = len(window) > max_chars
    result: dict = {
        "success": True,
        "path": str(path),
        "content": content,
        "offset": offset,
        "total_chars": total_chars,
    }
    if has_more:
        next_offset = offset + len(content)
        result["truncated"] = True
        result["next_offset"] = next_offset
        result["chars_remaining"] = total_chars - next_offset
    return json_result(**result)


def _write_file(args: dict, runtime: dict) -> str:
    path = resolve_workspace_path(args.get("path"))
    content = args.get("content")
    if content is None:
        return json_result(success=False, error="content is required")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(str(content), encoding="utf-8")
    return json_result(success=True, path=str(path), bytes=len(str(content).encode("utf-8")))


def _list_files(args: dict, runtime: dict) -> str:
    root = resolve_workspace_path(args.get("path") or ".")
    max_results = int(args.get("max_results") or 200)
    if root.is_file():
        return json_result(success=True, files=[str(root)])
    files: list[str] = []
    for item in root.rglob("*"):
        if ".git" in item.parts or "__pycache__" in item.parts:
            continue
        files.append(str(item.relative_to(resolve_workspace_path("."))))
        if len(files) >= max_results:
            break
    return json_result(success=True, root=str(root), files=files, truncated=len(files) >= max_results)


def _search_files(args: dict, runtime: dict) -> str:
    query = str(args.get("query") or "")
    if not query:
        return json_result(success=False, error="query is required")
    root = resolve_workspace_path(args.get("path") or ".")
    max_results = int(args.get("max_results") or 50)
    matches: list[dict[str, object]] = []
    for path in root.rglob("*"):
        if not path.is_file() or ".git" in path.parts or "__pycache__" in path.parts:
            continue
        try:
            for idx, line in enumerate(path.read_text(encoding="utf-8", errors="replace").splitlines(), 1):
                if query.lower() in line.lower():
                    matches.append({"path": str(path), "line": idx, "text": line[:300]})
                    break
        except OSError:
            continue
        if len(matches) >= max_results:
            break
    return json_result(success=True, matches=matches, truncated=len(matches) >= max_results)


def _patch_file(args: dict, runtime: dict) -> str:
    path = resolve_workspace_path(args.get("path"))
    old_text = args.get("old_text")
    new_text = args.get("new_text")
    if not old_text:
        return json_result(success=False, error="old_text is required")
    if new_text is None:
        return json_result(success=False, error="new_text is required")
    text = path.read_text(encoding="utf-8", errors="replace")
    count = text.count(str(old_text))
    if count != 1 and not args.get("replace_all"):
        return json_result(success=False, error=f"Expected one match, found {count}")
    updated = text.replace(str(old_text), str(new_text), -1 if args.get("replace_all") else 1)
    path.write_text(updated, encoding="utf-8")
    return json_result(success=True, path=str(path), replacements=count if args.get("replace_all") else 1)


registry.register(
    "read_file",
    {
        "description": (
            "Read a text file. Supports .txt, .md, .py, .json, .csv, .yaml, .xml, .html, .docx, and other UTF-8 text files. "
            "Absolute paths can also be read. "
            "Use offset+max_chars to paginate through large files: the result includes next_offset and chars_remaining "
            "so you know where to continue. Example: read_file(path=p, offset=0, max_chars=8000) then "
            "read_file(path=p, offset=8000, max_chars=8000). "
            "Large tool results are automatically cached to disk — use this tool with the given path to read them in chunks."
        ),
        "parameters": {
            "type": "object",
            "properties": {
                "path": {"type": "string"},
                "max_chars": {"type": "integer", "default": 20000, "description": "Max chars to return per call."},
                "offset": {"type": "integer", "default": 0, "description": "Start reading from this char position (for pagination)."},
            },
            "required": ["path"],
        },
    },
    _read_file,
)
registry.register(
    "write_file",
    {
        "description": "Write a UTF-8 text file. Use this to save reports or update project files.",
        "parameters": {
            "type": "object",
            "properties": {"path": {"type": "string"}, "content": {"type": "string"}},
            "required": ["path", "content"],
        },
    },
    _write_file,
)
registry.register(
    "list_files",
    {
        "description": "List files under a workspace path.",
        "parameters": {
            "type": "object",
            "properties": {"path": {"type": "string", "default": "."}, "max_results": {"type": "integer", "default": 200}},
            "required": [],
        },
    },
    _list_files,
)
registry.register(
    "search_files",
    {
        "description": "Search text files in the workspace for a query string.",
        "parameters": {
            "type": "object",
            "properties": {"query": {"type": "string"}, "path": {"type": "string", "default": "."}, "max_results": {"type": "integer", "default": 50}},
            "required": ["query"],
        },
    },
    _search_files,
)
registry.register(
    "patch_file",
    {
        "description": "Replace text in a file. Requires a unique old_text unless replace_all is true.",
        "parameters": {
            "type": "object",
            "properties": {
                "path": {"type": "string"},
                "old_text": {"type": "string"},
                "new_text": {"type": "string"},
                "replace_all": {"type": "boolean", "default": False},
            },
            "required": ["path", "old_text", "new_text"],
        },
    },
    _patch_file,
)
